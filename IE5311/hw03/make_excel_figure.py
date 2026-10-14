import numpy as np
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (CharacterProperties, Font as DFont,
                                   Paragraph, ParagraphProperties)
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, Side

import p1_relief as m

FONT = "Times New Roman"
NDRAW = 5000
SEED = 5311
MODELS = (("Risk neutral", False, False),
          ("Chance constrained", True, False),
          ("Chance and CVaR", True, True))

rng = np.random.default_rng(SEED)
draws = rng.choice(len(m.S), size=NDRAW, p=m.p)
counts = [int(np.sum(draws == s)) for s in m.S]

cost, orders = {}, []
for tag, ch, cv in MODELS:
    r = m.build(chance=ch, cvar=cv, name=tag)
    per = [m.recourse(r["x"], s) for s in m.S]
    cost[tag] = [sum(m.c[i] * r["x"][i] for i in m.I) + per[s][0]
                 for s in m.S]
    orders.append(tuple(sorted(m.S, key=lambda s: cost[tag][s])))

assert len(set(orders)) == 1, (
    "the three plans no longer rank the scenarios identically, so a single "
    "shared cumulative column is wrong; give each model its own")
ORDER = orders[0]

BODY = Font(name=FONT, size=11)
HEAD = Font(name=FONT, size=11, bold=True)
TITLE = Font(name=FONT, size=12, bold=True)
SUB = Font(name=FONT, size=12, italic=True)
RULE = Border(bottom=Side(style="thin", color="000000"))
MID = Alignment(horizontal="center")


def text(size):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(
        defRPr=CharacterProperties(latin=DFont(typeface=FONT), sz=size * 100,
                                   b=False)), endParaRPr=None)])


wb = Workbook()

ws = wb.active
ws.title = "Data"
ws["A1"] = "Table 1"
ws["A1"].font = TITLE
ws["A2"] = "Total Cost by Scenario Under Each First-Stage Plan"
ws["A2"].font = SUB

head = ["Scenario", "Probability", "Draws"] + [t for t, _, _ in MODELS]
for c, h in enumerate(head, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = HEAD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = RULE
ws.row_dimensions[4].height = 30

for s in m.S:
    r = 5 + s
    ws.cell(row=r, column=1, value=s + 1).font = BODY
    ws.cell(row=r, column=1).alignment = MID
    ws.cell(row=r, column=2, value=m.p[s]).font = BODY
    ws.cell(row=r, column=2).number_format = "0.00"
    ws.cell(row=r, column=2).alignment = MID
    ws.cell(row=r, column=3, value=counts[s]).font = BODY
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=3).alignment = MID
    for k, (tag, _, _) in enumerate(MODELS):
        cell = ws.cell(row=r, column=4 + k, value=cost[tag][s])
        cell.font = BODY
        cell.number_format = "#,##0.00"
for c in range(1, 7):
    ws.cell(row=15, column=c).border = Border(
        top=Side(style="thin", color="000000"))
for col, w in (("A", 10), ("B", 12), ("C", 10),
               ("D", 14), ("E", 14), ("F", 14)):
    ws.column_dimensions[col].width = w

cd = wb.create_sheet("Chart data")
cd["A1"] = "Cumulative proportion and plotting coordinates"
cd["A1"].font = TITLE

for c, h in enumerate(["Rank", "Scenario", "Cumulative"]
                      + [t for t, _, _ in MODELS], 1):
    cell = cd.cell(row=3, column=c, value=h)
    cell.font = HEAD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = RULE
cd.row_dimensions[3].height = 30

for i, s in enumerate(ORDER, 1):
    r = 3 + i
    cd.cell(row=r, column=1, value=i).font = BODY
    cd.cell(row=r, column=1).alignment = MID
    cd.cell(row=r, column=2, value=s + 1).font = BODY
    cd.cell(row=r, column=2).alignment = MID
    freq = f"Data!C{5 + s}/SUM(Data!$C$5:$C$14)"
    cd.cell(row=r, column=3,
            value=f"={freq}" if i == 1 else f"=C{r-1}+{freq}").font = BODY
    cd.cell(row=r, column=3).number_format = "0.0000"
    for k in range(3):
        cell = cd.cell(row=r, column=4 + k,
                       value=f"=Data!{chr(68 + k)}{5 + s}")
        cell.font = BODY
        cell.number_format = "#,##0.00"

for c, h in enumerate([t for t, _, _ in MODELS] + ["Cumulative"], 1):
    cell = cd.cell(row=16, column=c, value=h)
    cell.font = HEAD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = RULE
cd.row_dimensions[16].height = 30

for i in range(1, len(m.S) + 1):
    src = 3 + i
    for j, row in enumerate((15 + 2 * i, 16 + 2 * i)):
        for k in range(3):
            cell = cd.cell(row=row, column=1 + k,
                           value=f"={chr(68 + k)}{src}")
            cell.font = BODY
            cell.number_format = "#,##0.00"
        y = cd.cell(row=row, column=4)
        y.value = 0 if (i == 1 and j == 0) else \
            f"=C{src - 1 if j == 0 else src}"
        y.font = BODY
        y.number_format = "0.0000"
for col, w in (("A", 18), ("B", 18), ("C", 18), ("D", 14),
               ("E", 18), ("F", 18)):
    cd.column_dimensions[col].width = w

chart = ScatterChart()
chart.scatterStyle = "line"
chart.height = 9.5
chart.width = 16.0
chart.x_axis.title = "Total cost (dollars)"
chart.y_axis.title = "Cumulative proportion"
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.x_axis.scaling.min = 4500
chart.x_axis.scaling.max = 12500
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 1
chart.x_axis.numFmt = "#,##0"
chart.y_axis.numFmt = "0.0"
chart.x_axis.majorGridlines = None
chart.y_axis.majorGridlines = None
chart.x_axis.majorTickMark = "out"
chart.y_axis.majorTickMark = "out"
chart.x_axis.txPr = text(11)
chart.y_axis.txPr = text(11)
chart.x_axis.title.tx.rich.p[0].pPr = ParagraphProperties(
    defRPr=CharacterProperties(latin=DFont(typeface=FONT), sz=1200))
chart.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(
    defRPr=CharacterProperties(latin=DFont(typeface=FONT), sz=1200))
chart.legend.position = "b"
chart.legend.txPr = text(11)
chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))

y = Reference(cd, min_col=4, min_row=17, max_row=36)
for k, (tag, _, _) in enumerate(MODELS):
    ser = Series(y, Reference(cd, min_col=1 + k, min_row=17, max_row=36),
                 title=tag)
    ser.marker.symbol = "none"
    ser.smooth = False
    ser.graphicalProperties.line.solidFill = "000000"
    ser.graphicalProperties.line.width = 15875
    ser.graphicalProperties.line.dashStyle = ("solid", "dash", "sysDot")[k]
    chart.series.append(ser)

wb.create_chartsheet("Figure").add_chart(chart)
wb.move_sheet("Figure", offset=-2)

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True

wb.save("p1_cost_distributions.xlsx")
print("wrote p1_cost_distributions.xlsx")
print("sheets:", wb.sheetnames)
